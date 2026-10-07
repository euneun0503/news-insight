#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
기존 Colab 노트북이 만든 엑셀 파일을 새 데이터 저장소로 가져오기 (과거 자료 이관용)

  python collector/import_excel.py 엑셀1.xlsx 엑셀2.xlsx ...

지원 형식
  - 랭킹 파일  (네이버_랭킹뉴스_*.xlsx) : 매체명 시트 [날짜, 제목, 조회수, 링크]
  - 발행 파일  (*_상세 시트)            : [발행일, 발행시각, 제목, 링크]
같은 기사(oid/aid)는 자동으로 병합되므로 여러 번 가져와도 중복되지 않습니다.
"""
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import store  # noqa: E402
from store import CONFIG, article_key, clean_text, now_kst  # noqa: E402

NAME2OID = {v: k for k, v in CONFIG["media"].items()}


def norm_day(v):
    s = str(v).strip()[:10].replace(".", "-").replace("/", "-")
    if len(s) == 8 and s.isdigit():
        s = f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return s if len(s) == 10 and s[4] == "-" else None


def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    for r in it:
        yield dict(zip(header, r))


def import_file(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ranking = defaultdict(lambda: defaultdict(list))  # day -> oid -> items
    articles = defaultdict(lambda: defaultdict(list))
    for ws in wb.worksheets:
        name = ws.title
        if name.endswith("_상세"):
            oid = NAME2OID.get(name[:-3])
            if not oid:
                continue
            for r in rows(ws):
                key = article_key(str(r.get("링크") or ""))
                day = norm_day(r.get("발행일") or "")
                if not key or not day:
                    continue
                hm = str(r.get("발행시각") or "")[:5]
                articles[day][oid].append((key[0], key[1], hm, clean_text(r.get("제목")), clean_text(r.get("기자"))))
        elif name in NAME2OID:
            oid = NAME2OID[name]
            for r in rows(ws):
                key = article_key(str(r.get("링크") or ""))
                day = norm_day(r.get("날짜") or "")
                if not key or not day:
                    continue
                lst = ranking[day][oid]
                v = r.get("조회수")
                views = None if v in (None, "", "미공개") else int(v)
                lst.append([key[0], key[1], len(lst) + 1, views, clean_text(r.get("제목")),
                            clean_text(r.get("기자")), str(r.get("입력시각") or "")[:16]])
    # 한 매체의 하루 조회수가 전부 0 이면 '조회수 미제공 매체'로 처리 (실제 0회가 아님)
    for per in ranking.values():
        for items in per.values():
            if items and all(not r[3] for r in items):  # 전부 0 또는 빈칸
                for r in items:
                    r[3] = None
    ts = now_kst().isoformat(timespec="seconds")
    for day, per in ranking.items():
        store.save_ranking(day, dict(per), ts, final=True)
    for day, per in articles.items():
        for oid, items in per.items():
            store.merge_articles(day, oid, items, ts)
    print(f"  {Path(path).name}: 랭킹 {len(ranking)}일, 발행 {len(articles)}일")
    return set(ranking) | set(articles)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return
    touched = set()
    for p in sys.argv[1:]:
        touched |= import_file(p)
    store.rebuild(touched)
    store.build_meta()
    print(f"✅ 가져오기 완료 ({len(touched)}일)")


if __name__ == "__main__":
    main()
